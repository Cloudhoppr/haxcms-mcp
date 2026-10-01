#!/usr/bin/env python
"""Regenerate the Phase 7 import fixtures (PLAN T7.6).

    uv run python tests/fixtures/import/generate.py

The binaries are COMMITTED next to this script so the suites never need the generator
libraries at test time. Timestamps are pinned where the libraries allow it, but the ZIP
container stamps still vary between runs — fixtures are validated by the live suites, not
by byte-diffing.

Fixture contracts (source-verified against ../haxcms-nodejs):

* syllabus.docx — mammoth maps Word Heading 1/2/3 styles to h1/h2/h3, so the outline
  shape is: two H1 roots ("Biology 101 Syllabus" with two H2 children, "Appendix: Lab
  Safety" without); the H3 folds into its H2's content stream (importHtmlToItems only
  nests childTag = h(level+1)). The inline image is a real PNG (Songline_1) so the
  createSite path can materialize it as a site file.
* slides.pptx — three titled slides (pptx-in-html-out emits one section per slide).
* workbook.xlsx — importXlsx reads the FIRST sheet and requires a header row with
  `title,slug,parent,content` columns (getHeaderLookup); title and slug are mandatory
  per row and slugs must be unique. The second sheet exercises convertXlsxToCsv's
  ?sheet= selection.
* onepage.pdf — one page of text for import-pdf / pdfToHtml.
* page.html — h1/h2 structure for the html document import and the html PLATFORM
  multipart import (the repoUrl mode is unreachable for local files: the safeFetch SSRF
  guard refuses loopback addresses).
"""

from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent
IMAGES = HERE.parent / "images"
SONGLINE_1 = IMAGES / "Songline_1.png"  # 120 x 80 real png
EPOCH = datetime(2026, 1, 1, 0, 0, 0, tzinfo=UTC)


def _pin_docx(document: object) -> None:
    props = document.core_properties  # type: ignore[attr-defined]
    props.created = EPOCH
    props.modified = EPOCH
    props.author = "HAXcms MCP fixtures"
    props.last_modified_by = "HAXcms MCP fixtures"
    props.revision = 1


def make_docx(path: Path) -> None:
    from docx import Document
    from docx.shared import Inches

    document = Document()
    document.add_heading("Biology 101 Syllabus", level=1)
    document.add_paragraph(
        "Welcome to Biology 101. This course covers cells, genetics and evolution across two units."
    )
    document.add_heading("Unit One: Cells", level=2)
    document.add_paragraph("Cell structure and function, from membranes to mitochondria.")
    document.add_picture(str(SONGLINE_1), width=Inches(1.2))
    document.add_heading("Unit Two: Genetics", level=2)
    document.add_paragraph("Mendelian inheritance and the machinery of heredity.")
    document.add_heading("Lesson: Punnett Squares", level=3)
    document.add_paragraph("Predicting genotype ratios with a Punnett square.")
    document.add_heading("Appendix: Lab Safety", level=1)
    document.add_paragraph("Goggles on, no open-toe shoes, report every spill.")
    _pin_docx(document)
    document.save(str(path))


def make_pptx(path: Path) -> None:
    from pptx import Presentation

    deck = Presentation()
    layout = deck.slide_layouts[1]  # Title and Content
    bodies = [
        ("Introduction", "What this study measures and why it matters."),
        ("Methods", "Sample preparation, controls and the measurement pipeline."),
        ("Results", "Three findings, each with its confidence interval."),
    ]
    for title, body in bodies:
        slide = deck.slides.add_slide(layout)
        slide.shapes.title.text = title  # type: ignore[union-attr]
        slide.placeholders[1].text_frame.text = body
    props = deck.core_properties
    props.created = EPOCH
    props.modified = EPOCH
    props.author = "HAXcms MCP fixtures"
    props.last_modified_by = "HAXcms MCP fixtures"
    props.revision = 1
    deck.save(str(path))


def make_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    schedule = workbook.active
    schedule.title = "Schedule"
    schedule.append(["title", "slug", "parent", "content"])
    schedule.append(["Course Home", "course-home", "", "<p>Weekly schedule for the course.</p>"])
    schedule.append(["Week One", "week-one", "course-home", "<p>Cells.</p>"])
    schedule.append(["Week Two", "week-two", "course-home", "<p>Genetics.</p>"])
    roster = workbook.create_sheet("Roster")
    roster.append(["title", "slug", "parent", "content"])
    roster.append(["Ada", "ada", "", "<p>Student.</p>"])
    roster.append(["Grace", "grace", "", "<p>Student.</p>"])
    props = workbook.properties
    props.created = EPOCH
    props.modified = EPOCH
    props.creator = "HAXcms MCP fixtures"
    workbook.save(str(path))


def make_pdf(path: Path) -> None:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.creation_date = EPOCH
    pdf.set_title("Quantum Field Theory Notes")
    pdf.set_author("HAXcms MCP fixtures")
    pdf.add_page()
    pdf.set_font("Helvetica", size=18)
    pdf.cell(text="Quantum Field Theory Notes", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=12)
    pdf.multi_cell(
        w=0,
        text=(
            "These notes cover the path integral formulation of quantum field theory. "
            "The first chapter reviews canonical quantization of the free scalar field."
        ),
    )
    pdf.output(str(path))


PAGE_HTML = """<!doctype html>
<html lang="en">
  <head>
    <meta charset="utf-8" />
    <title>Field Trip</title>
  </head>
  <body>
    <h1>Field Trip</h1>
    <p>Everything the class needs to know about the field trip.</p>
    <h2>Logistics</h2>
    <p>The bus leaves at 8am sharp from the north parking lot.</p>
    <h2>Packing List</h2>
    <p>Water, lunch, sunscreen and a notebook.</p>
  </body>
</html>
"""


def main() -> None:
    make_docx(HERE / "syllabus.docx")
    make_pptx(HERE / "slides.pptx")
    make_xlsx(HERE / "workbook.xlsx")
    make_pdf(HERE / "onepage.pdf")
    (HERE / "page.html").write_text(PAGE_HTML, encoding="utf-8", newline="\n")
    print("wrote syllabus.docx, slides.pptx, workbook.xlsx, onepage.pdf, page.html")


if __name__ == "__main__":
    main()
